#!/usr/bin/env python3
"""Build the revised manuscript and response-to-comments .docx files.

Converts the markdown masters in Revision_2026/manuscript/ to Word documents,
inserting the data tables from Revision_2026/tables/ after their legends.

Usage:
    python3 12_build_manuscript_docx.py
"""
import csv
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MSDIR = os.path.join(BASE, "manuscript")
TABDIR = os.path.join(BASE, "tables")

# legend prefix -> (tsv file, column renaming map, columns to drop)
TABLE_MAP = {
    "Table 1.": ("Table1_assembly_sequencing.tsv", {}, []),
    "Table 2.": ("Table_TE_classes_core_vs_accessory.tsv", {}, []),
    "Table 3.": ("Table_centromeres_gc.tsv",
                 {"chrom": "Chromosome", "compartment": "Compartment",
                  "start": "Start (bp)", "end": "End (bp)",
                  "size_kb": "Size (kb)", "gc": "GC at trough",
                  "gc_chrom_median": "Chromosome median GC",
                  "gc_depression": "GC depression",
                  "gc_pctile": "GC pctile",
                  "n_window_pairs": "Window pairs",
                  "median_identity": "Median identity (%)",
                  "gene_density_pctile": "Gene density pctile",
                  "TE_density_pctile": "TE density pctile",
                  "support": "Support (0-3)"},
                 ["peak", "peak_bp", "relative_position", "n_gc_troughs",
                  "pct_band_pairs_in_call", "band_pairs_per_mb",
                  "assembly_gap_in_call", "shift_vs_identity_only_kb",
                  "overlap_identity_only_pct"]),
    "Table 4.": ("Table_per_chromosome_content.tsv",
                 {"chrom": "Chromosome", "compartment": "Compartment",
                  "length_bp": "Length (bp)", "genes": "Genes",
                  "genes_per_Mb": "Genes per Mb",
                  "exons_per_gene": "Exons per gene", "CDS_bp": "CDS (bp)",
                  "coding_pct": "Coding (%)", "repeat_bp": "Repeat (bp)",
                  "repeat_pct": "Repeat (%)",
                  "intact_TE_elements": "Intact TE elements",
                  "BGCs": "BGCs", "effectors": "Effector candidates",
                  "effectors_per_Mb": "Effectors per Mb",
                  "unique_effectors": "Sequence-unique effectors"}, []),
    "Table 5.": ("Table5_three_strain.tsv", {}, []),
    "Table 6.": ("Table_FOSC_phylogeny_taxa.tsv",
                 {"tip": "Tip label", "genomeID": "Genome ID",
                  "accession": "Assembly accession", "organism": "Organism"}, []),
}


def build_table5():
    """Compose Table 5 (three-strain effector comparison + enrichment tests)."""
    rows = [["Section", "Category", "Clusters / a", "MRL8996 / b",
             "Fol4287 / c", "Fo47 / d", "Odds ratio", "p"]]
    labels = {"core_all_three": "Shared by all three strains",
              "MRL8996_Fol4287": "MRL8996 and Fol4287 only",
              "MRL8996_Fo47": "MRL8996 and Fo47 only",
              "Fol4287_Fo47": "Fol4287 and Fo47 only",
              "MRL8996_unique": "MRL8996 only (sequence-unique)",
              "Fol4287_unique": "Fol4287 only",
              "Fo47_unique": "Fo47 only",
              "TOTAL": "Total"}
    with open(os.path.join(TABDIR, "venn_counts.tsv")) as fh:
        for r in csv.DictReader(fh, delimiter="\t"):
            rows.append(["Cluster occupancy", labels.get(r["Region"], r["Region"]),
                         r["n_clusters"], r["n_MRL8996_proteins"],
                         r["n_Fol4287_proteins"], r["n_Fo47_proteins"], "", ""])
    with open(os.path.join(TABDIR, "unique_effector_enrichment.tsv")) as fh:
        for r in csv.DictReader(fh, delimiter="\t"):
            rows.append(["Fisher exact test", r["Test"].replace("_", " "),
                         r["a"], r["b"], r["c"], r["d"],
                         r["odds_ratio"], "%.3f" % float(r["p_value"])])
    out = os.path.join(TABDIR, "Table5_three_strain.tsv")
    with open(out, "w", newline="") as fh:
        csv.writer(fh, delimiter="\t").writerows(rows)


def add_tsv_table(doc, fname, rename, drop):
    path = os.path.join(TABDIR, fname)
    with open(path) as fh:
        rows = [r for r in csv.reader(fh, delimiter="\t")]
    header = [rename.get(c, c.replace("_", " ")) for c in rows[0]]
    keep = [i for i, c in enumerate(rows[0]) if c not in drop]
    tbl = doc.add_table(rows=1, cols=len(keep))
    tbl.style = "Table Grid"
    # keep wide tables inside the A4 text block (17 cm at 2 cm margins)
    if len(keep) >= 6:
        first = Cm(5.0)
        rest = Cm((17.0 - 5.0) / (len(keep) - 1))
        tbl.autofit = False
        widths = [first] + [rest] * (len(keep) - 1)
        for row in tbl.rows:
            for c, w in zip(row.cells, widths):
                c.width = w
        tbl._widths = widths
    else:
        tbl._widths = None
    for j, i in enumerate(keep):
        cell = tbl.rows[0].cells[j]
        cell.text = header[i]
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    tr = tbl.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    trPr.append(hdr)
    for r in rows[1:]:
        cells = tbl.add_row().cells
        if tbl._widths:
            for c, w in zip(cells, tbl._widths):
                c.width = w
        for j, i in enumerate(keep):
            val = r[i] if i < len(r) else ""
            cells[j].text = "" if val in ("nan", "") else val
            for run in cells[j].paragraphs[0].runs:
                run.font.size = Pt(8)
    doc.add_paragraph("")


def emit_runs(par, text):
    """Render **bold** and *italic* spans."""
    for piece in re.split(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def md_to_docx(md_path, docx_path, insert_tables=False):
    if insert_tables:
        build_table5()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(2.0)

    lines = open(md_path).read().split("\n")
    md_table = []
    for raw in lines:
        line = raw.rstrip()

        # markdown pipe table (response letter summary)
        if line.startswith("|"):
            md_table.append([c.strip() for c in line.strip("|").split("|")])
            continue
        if md_table:
            body = [r for r in md_table if not set("".join(r)) <= set("-: ")]
            tbl = doc.add_table(rows=0, cols=len(body[0]))
            tbl.style = "Table Grid"
            for k, r in enumerate(body):
                cells = tbl.add_row().cells
                for j, v in enumerate(r[:len(body[0])]):
                    cells[j].text = v
                    for run in cells[j].paragraphs[0].runs:
                        run.font.size = Pt(9)
                        run.bold = (k == 0)
            doc.add_paragraph("")
            md_table = []

        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("---"):
            doc.add_paragraph("")
        elif re.match(r"^\d+\. ", line):
            emit_runs(doc.add_paragraph(style="List Number"), line.split(". ", 1)[1])
        else:
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(6)
            emit_runs(par, line)
            if insert_tables:
                m = re.match(r"^\*\*(Table \d\.)", line)
                if m and m.group(1) in TABLE_MAP:
                    fname, rename, drop = TABLE_MAP[m.group(1)]
                    add_tsv_table(doc, fname, rename, drop)
    doc.save(docx_path)
    print("wrote %s" % docx_path)


def main():
    md_to_docx(os.path.join(MSDIR, "ms_v2.0_revised.md"),
               os.path.join(MSDIR, "MRL8996_2026_v2.0_revised.docx"),
               insert_tables=True)
    md_to_docx(os.path.join(MSDIR, "response_to_comments.md"),
               os.path.join(MSDIR, "Response_to_comments.docx"))


if __name__ == "__main__":
    main()
